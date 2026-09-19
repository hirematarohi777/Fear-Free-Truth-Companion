import os
import json
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import docx

OUTPUT_DOCS_DIR = os.path.join(os.path.dirname(__file__), "documents")
OUTPUT_WEB_DIR = os.path.join(os.path.dirname(__file__), "websites")
OUTPUT_EXP_DIR = os.path.join(os.path.dirname(__file__), "expected_results")

os.makedirs(OUTPUT_DOCS_DIR, exist_ok=True)
os.makedirs(OUTPUT_WEB_DIR, exist_ok=True)
os.makedirs(OUTPUT_EXP_DIR, exist_ok=True)


def build_pdf(filename: str, title: str, paragraphs: list):
    path = os.path.join(OUTPUT_DOCS_DIR, filename)
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=50, rightMargin=50, topMargin=50, bottomMargin=50)
    styles = getSampleStyleSheet()
    title_style = styles['Heading1']
    body_style = styles['Normal']
    
    story = [Paragraph(f"<b>{title}</b>", title_style), Spacer(1, 15)]
    for p in paragraphs:
        story.append(Paragraph(p, body_style))
        story.append(Spacer(1, 10))
    doc.build(story)
    print(f"Generated PDF: {path}")


def build_docx(filename: str, title: str, sections: list):
    path = os.path.join(OUTPUT_DOCS_DIR, filename)
    doc = docx.Document()
    doc.add_heading(title, 0)
    for heading, text in sections:
        if heading:
            doc.add_heading(heading, level=1)
        doc.add_paragraph(text)
    doc.save(path)
    print(f"Generated DOCX: {path}")


def generate_all():
    # 1. Standard Home Loan (PDF)
    build_pdf(
        "1_standard_home_loan.pdf",
        "SANCTION LETTER - BENGALURU HOUSING FINANCE LTD",
        [
            "Borrower Name: Smt. Priya Sharma & Sri. Rajesh Sharma",
            "Sanctioned Loan Amount: Rs. 40,00,000 (Rupees Forty Lakhs Only)",
            "Facility: Regular Home Loan (Floating Rate)",
            "Rate of Interest: 8.5% p.a. (linked to External Benchmark Lending Rate)",
            "Repayment Period: 240 months (20 Years)",
            "Monthly EMI: Rs. 34,713.00 payable via NACH mandate.",
            "Processing fee: 0.5% of sanctioned loan amount payable upfront prior to documentation.",
            "Tax Treatment: All fees are subject to 18% Goods and Services Tax (GST) as applicable under Indian tax laws.",
            "Prepayment terms: Nil foreclosure or prepayment charges for individual borrowers on floating rate.",
            "Special Note: Synthetic demonstration data — not a real lender assessment."
        ]
    )

    # 2. Additional Charges Example (DOCX)
    build_docx(
        "2_additional_charges_loan.docx",
        "LOAN OFFER LETTER - METRO CREDIT SERVICES",
        [
            ("1. Sanction Summary", "Sanctioned Loan Amount: Rs. 50,00,000. Rate of Interest: 8.9% p.a. Tenor: 180 months."),
            ("2. Disclosed & Administrative Fees", "Processing charges: Rs. 25,000. Administrative charges: Rs. 12,500 for document custody and legal verification."),
            ("3. Inspection and Technical Valuation", "Legal & valuation charges: Rs. 7,500 per site inspection."),
            ("4. Insurance Coverage", "Property and Credit Shield Insurance Cover: Rs. 42,000 mandatory protection bundled into loan sanction."),
            ("5. Delayed Payment & Default", "Late payment charges: Rs. 500 per month plus penal interest of 2% per month on overdue installments. Cheque bounce: Rs. 450 per instance."),
            ("6. Disclaimer", "Synthetic demonstration data — not a real lender assessment.")
        ]
    )

    # 3. Conflicting Document (PDF)
    build_pdf(
        "3_conflicting_terms_loan.pdf",
        "SPECIAL HOUSING FACILITY AGREEMENT",
        [
            "Borrower: Sri. Arjun Kumar",
            "Sanctioned Loan Amount: Rs. 35,00,000. Tenor: 180 months.",
            "Section 2 (Schedule of Rates): Rate of Interest: 8.75% p.a. floating rest.",
            "Section 8 (General Terms of Repayment): Subject to revised guidelines, Rate of Interest: 9.25% p.a. effective upon second disbursement tranche.",
            "Section 11 (Upfront Levies): Processing fee is subject to prevailing tariff schedule without stated percentage or cap.",
            "Section 14 (Early Closure): Prepayment is permitted subject to conditions outlined in branch circulars.",
            "Notice: Synthetic demonstration data — not a real lender assessment."
        ]
    )

    # 4. Scanned Degraded Document (PDF)
    build_pdf(
        "4_scanned_degraded_document.pdf",
        "SANCTION SUMMARY (SCANNED ARCHIVE)",
        [
            "Borrower: M/s Vikram Rao",
            "Sanctioned Loan Amount: Rs. 28,00,000. Rate of Interest: 9.1% p.a.",
            "[Note: Subordinate clause text faded or low contrast scan requiring OCR review]",
            "Processing fee: 0.75%",
            "Synthetic demonstration data — not a real lender assessment."
        ]
    )

    # 5. Suspicious Website Fixture (JSON)
    suspicious_web = {
        "url": "http://fast-instant-sanction-now.test/apply",
        "page_title": "Instant 100% Guaranteed Loan - No CIBIL Required",
        "page_text": (
            "Welcome to Fast Sanction! Guaranteed 100% loan approval within 5 minutes without any credit check or CIBIL score! "
            "Please pay processing fee in advance before disbursement of Rs. 4,999 to our executive account. "
            "Enter your netbanking password to verify KYC."
        ),
        "claimed_entity": "State Bank of India",
        "intended_outcome": "Significant risk indicators found",
        "expected_flags": [
            "Transport Layer Security (HTTPS) - CONCERN (uses HTTP)",
            "Unrealistic Financial & Approval Claims - CONCERN (100% guarantee)",
            "Advance Fee Payment Requirement - CONCERN (pay in advance)",
            "Sensitive Credential Solicitation - CONCERN (netbanking password)",
            "Official Regulatory & Domain Registry Verification - CONCERN (Claims SBI but uses .test domain)"
        ]
    }
    with open(os.path.join(OUTPUT_WEB_DIR, "5_suspicious_website.json"), "w", encoding="utf-8") as f:
        json.dump(suspicious_web, f, indent=2)

    # 6. Incomplete Website Fixture (JSON)
    incomplete_web = {
        "url": "https://urban-fintech-cooperative.test",
        "page_title": "Urban Fintech Cooperative Services",
        "page_text": (
            "Urban Fintech Cooperative offers community mortgage solutions at standard rates. "
            "We act as an aggregator connecting borrowers with lending partners. "
            "Contact our representative for details. Privacy policy and grievance officer details under review."
        ),
        "claimed_entity": None,
        "intended_outcome": "Some concerns require verification",
        "expected_flags": [
            "Grievance Redressal & Privacy Disclosure - CONCERN (Missing nodal contact)",
            "Official Regulatory & Domain Registry Verification - INCONCLUSIVE (No recognized SCB/HFC)"
        ]
    }
    with open(os.path.join(OUTPUT_WEB_DIR, "6_incomplete_website.json"), "w", encoding="utf-8") as f:
        json.dump(incomplete_web, f, indent=2)

    expected = {
        "1_standard_home_loan": {
            "loanAmount": 4000000,
            "interestRatePercent": 8.5,
            "tenureMonths": 240,
            "statedEmi": 34713,
            "upfrontProcessingFeePercent": 0.5,
        },
        "2_additional_charges_loan": {
            "loanAmount": 5000000,
            "interestRatePercent": 8.9,
            "tenureMonths": 180,
            "expectedFindings": [
                "Administrative charges",
                "Legal & valuation",
                "Insurance",
                "Late payment",
                "Cheque bounce",
            ],
        },
        "3_conflicting_terms_loan": {
            "loanAmount": 3500000,
            "conflictingRates": [8.75, 9.25],
            "processingFee": "not_found",
        },
        "4_scanned_degraded_document": {
            "loanAmount": 2800000,
            "interestRatePercent": 9.1,
            "ocrLimitationsExpected": True,
        },
        "5_suspicious_website": suspicious_web,
        "6_incomplete_website": incomplete_web,
    }
    with open(os.path.join(OUTPUT_EXP_DIR, "baseline_validation.json"), "w", encoding="utf-8") as f:
        json.dump(expected, f, indent=2)

    print("All synthetic fixtures generated successfully!")


if __name__ == "__main__":
    generate_all()
